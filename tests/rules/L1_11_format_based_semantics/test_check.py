"""L1-11: 書式によるデータ区別チェックのテスト"""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from harunobu.core.analyzer import Analyzer
from harunobu.core.models import CellFormat
from harunobu.rules.level1.L1_11_format_based_semantics import (
    FormatBasedSemanticsRule,
)


class TestFormatBasedSemanticsCheck:
    """L1-11 書式依存データ区別チェック"""

    @pytest.fixture
    def rule(self):
        return FormatBasedSemanticsRule()

    def test_no_format_issues_passes(self, rule, create_context):
        """書式の区別がない場合はpass"""
        ctx = create_context({"Sheet1": [["名前", "年齢"], ["太郎", 20], ["花子", 25]]})
        result = rule.check(ctx)
        assert result.passed is True
        # assert result.score == 100
        assert len(result.violations) == 0

    def test_csv_skipped(self, rule, create_context):
        """CSV形式の場合はチェック対象外"""
        ctx = create_context(
            {"Sheet1": [["名前", "年齢"], ["太郎", 20]]},
            file_name="test.csv",
        )
        result = rule.check(ctx)
        assert result.passed is True

    def test_varied_font_colors_detected(self, rule, create_context):
        """列内でフォント色が複数使われている場合は違反"""
        ctx = create_context(
            {
                "Sheet1": [
                    ["ステータス", "値"],
                    ["正常", 100],
                    ["異常", 50],
                    ["正常", 80],
                ]
            }
        )
        # body行のセルにフォント色を設定
        ctx.sheet.cells[(2, 1)].fmt = CellFormat(font_color="000000")
        ctx.sheet.cells[(3, 1)].fmt = CellFormat(font_color="FF0000")  # 赤
        ctx.sheet.cells[(4, 1)].fmt = CellFormat(font_color="000000")
        result = rule.check(ctx)
        assert result.passed is False
        assert len(result.violations) >= 1

    def test_partial_bold_detected(self, rule, create_context):
        """一部のセルのみ太字の場合は違反"""
        ctx = create_context(
            {
                "Sheet1": [
                    ["名前", "点数"],
                    ["太郎", 100],
                    ["花子", 50],
                    ["次郎", 80],
                ]
            }
        )
        # 1セルだけ太字
        ctx.sheet.cells[(2, 2)].fmt = CellFormat(bold=True)
        result = rule.check(ctx)
        assert result.passed is False
        assert any("太字" in v.description for v in result.violations)

    def test_partial_italic_detected(self, rule, create_context):
        """一部のセルのみ斜体の場合は違反"""
        ctx = create_context(
            {
                "Sheet1": [
                    ["名前", "点数"],
                    ["太郎", 100],
                    ["花子", 50],
                    ["次郎", 80],
                ]
            }
        )
        ctx.sheet.cells[(3, 2)].fmt = CellFormat(italic=True)
        result = rule.check(ctx)
        assert result.passed is False
        assert any("斜体" in v.description for v in result.violations)

    def test_varied_bg_colors_detected(self, rule, create_context):
        """背景色が複数の非デフォルト色で使用されている場合は違反"""
        ctx = create_context(
            {
                "Sheet1": [
                    ["ステータス", "値"],
                    ["正常", 100],
                    ["異常", 50],
                    ["警告", 80],
                ]
            }
        )
        ctx.sheet.cells[(2, 1)].fmt = CellFormat(bg_color="00FF00")  # 緑
        ctx.sheet.cells[(3, 1)].fmt = CellFormat(bg_color="FF0000")  # 赤
        ctx.sheet.cells[(4, 1)].fmt = CellFormat(bg_color="FFFF00")  # 黄
        result = rule.check(ctx)
        assert result.passed is False
        assert any("背景色" in v.description for v in result.violations)

    @pytest.mark.parametrize(
        ("colors", "expected_passed"),
        [
            pytest.param([None, None, None], True, id="unfilled"),
            pytest.param(["FFFF00", "FFFF00", "FFFF00"], True, id="uniform-yellow"),
            pytest.param(["FFFF00", "00FFFF00", "FFFFFF00"], True, id="same-rgb-different-alpha"),
            pytest.param([None, "FFFFFF", "00FFFFFF", "FFFFFFFF"], True, id="unfilled-and-white"),
            pytest.param([None, "FFFFFF00", None], False, id="unfilled-and-yellow"),
            pytest.param(["FFFFFF", "00FFFF00", None], False, id="white-and-yellow"),
            pytest.param([None, "00000000", None], False, id="unfilled-and-black"),
            pytest.param(["00FF00", "FF0000", "FF0000"], False, id="two-colors"),
            pytest.param([None, "FFFF00", None, "FFFF00"], False, id="manual-stripes"),
        ],
    )
    def test_background_color_variations(self, rule, create_context, colors, expected_passed):
        rows = [["値"]] + [[str(i)] for i in range(len(colors))]
        ctx = create_context({"Sheet1": rows})
        for row, color in enumerate(colors, start=2):
            ctx.sheet.cells[(row, 1)].fmt = CellFormat(bg_color=color)

        result = rule.check(ctx)

        assert result.passed is expected_passed
        if not expected_passed:
            assert any(v.cell_range == "列1" and "背景色" in v.description for v in result.violations)

    def test_header_and_empty_cell_backgrounds_are_ignored(self, rule, create_context):
        ctx = create_context({"Sheet1": [["値"], [1], [None], [2]]})
        ctx.sheet.cells[(1, 1)].fmt = CellFormat(bg_color="FFFF00")
        ctx.sheet.cells[(3, 1)].fmt = CellFormat(bg_color="FFFF00")

        assert rule.check(ctx).passed is True

    def test_all_bold_passes(self, rule, create_context):
        """全セルが太字の場合は違反にならない"""
        ctx = create_context(
            {
                "Sheet1": [
                    ["名前", "点数"],
                    ["太郎", 100],
                    ["花子", 50],
                ]
            }
        )
        ctx.sheet.cells[(2, 2)].fmt = CellFormat(bold=True)
        ctx.sheet.cells[(3, 2)].fmt = CellFormat(bold=True)
        result = rule.check(ctx)
        # 全セル太字 → bold_count == total_cells → 違反なし
        assert not any("太字" in v.description for v in result.violations)

    def test_single_cell_column_skipped(self, rule, create_context):
        """データが1セルのみの列はチェック対象外"""
        ctx = create_context(
            {
                "Sheet1": [
                    ["名前", "値"],
                    ["太郎", 100],
                ]
            }
        )
        ctx.sheet.cells[(2, 1)].fmt = CellFormat(font_color="FF0000")
        result = rule.check(ctx)
        # total_cells < 2 → スキップ
        assert result.passed is True


def test_unfilled_and_yellow_rows_are_detected(monkeypatch):
    monkeypatch.setenv("HARUNOBU_AI_DISABLED", "true")
    workbook = Workbook()
    sheet = workbook.active
    for row in [
        ("支店", "売上"),
        ("東京", 1200),
        ("大阪", 800),
        ("福岡", 900),
        ("札幌", 600),
    ]:
        sheet.append(row)
    for cell in sheet[3]:
        cell.fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")

    data = BytesIO()
    workbook.save(data)
    data.seek(0)

    result = Analyzer(mode="lite").analyze_file("sales.xlsx", data)
    check = result.sheets[0].tables[0].mr_result.all_results()["L1-11"]

    assert check.passed is False
    assert any("背景色" in violation.description for violation in check.violations)


@pytest.mark.parametrize("two_pass", [False, True], ids=["single-pass", "two-pass"])
@pytest.mark.parametrize("highlight", [False, True], ids=["table-style-only", "manual-yellow"])
def test_excel_table_style_and_manual_fill(monkeypatch, two_pass, highlight):
    monkeypatch.setenv("HARUNOBU_AI_DISABLED", "true")
    workbook = Workbook()
    sheet = workbook.active
    for row in [("支店", "売上"), ("東京", 1200), ("大阪", 800), ("福岡", 900), ("札幌", 600)]:
        sheet.append(row)
    table = Table(displayName="Sales", ref="A1:B5")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    sheet.add_table(table)
    if highlight:
        for cell in sheet[3]:
            cell.fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")

    data = BytesIO()
    workbook.save(data)
    data.seek(0)
    if two_pass:
        monkeypatch.setattr("harunobu.core.reader._SINGLE_PASS_THRESHOLD", 0)

    result = Analyzer(mode="lite").analyze_file("sales.xlsx", data)
    check = result.sheets[0].tables[0].mr_result.all_results()["L1-11"]

    assert check.passed is not highlight
    assert any("背景色" in v.description for v in check.violations) is highlight

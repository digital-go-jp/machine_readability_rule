"""L1-13 用サンプル（非表示の行・列）。

OK: 非表示なしの人口表。
NG-01: テーブル内のデータ行を非表示にした例。
NG-02: テーブル内のデータ列を非表示にした例。
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .common import write_base_data


def _write_ok(ws) -> None:
    write_base_data(ws)


def _write_ng01(ws) -> None:
    """データ行の 1 行（例: 宮城県行 = 5 行目）を非表示。"""
    write_base_data(ws)
    ws.row_dimensions[5].hidden = True


def _write_ng02(ws) -> None:
    """年度列の 1 列（例: 2022 年 = 列 D）を非表示。"""
    write_base_data(ws)
    ws.column_dimensions[get_column_letter(4)].hidden = True


def generate(output_dir: Path) -> list[Path]:
    path = output_dir / "rule_13_hidden_rows_columns.xlsx"
    wb = Workbook()

    ws_ok = wb.active
    ws_ok.title = "OK"
    _write_ok(ws_ok)

    ws_ng1 = wb.create_sheet("NG-01")
    _write_ng01(ws_ng1)

    ws_ng2 = wb.create_sheet("NG-02")
    _write_ng02(ws_ng2)

    wb.save(path)
    return [path]

"""サンプル生成用の共通処理。"""

from __future__ import annotations

import csv
from copy import copy
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

PREFECTURES = ["北海道", "青森県", "岩手県", "宮城県", "秋田県"]
YEARS = [2020, 2021, 2022, 2023, 2024]
POP_DATA = [
    [5250, 5224, 5188, 5140, 5092],
    [1246, 1228, 1211, 1193, 1175],
    [1227, 1211, 1195, 1178, 1162],
    [2302, 2292, 2280, 2266, 2252],
    [960, 941, 923, 905, 888],
]


def make_workbook() -> tuple[Workbook, object, object]:
    workbook = Workbook()
    ok_sheet = workbook.active
    ok_sheet.title = "OK"
    ng_sheet = workbook.create_sheet("NG")
    return workbook, ok_sheet, ng_sheet


def make_empty_workbook(*, first_sheet_title: str = "Sheet1") -> tuple[Workbook, object]:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = first_sheet_title
    return workbook, sheet


def write_base_data(ws, *, title_row: bool = False, start_row: int = 1) -> int:
    """人口データの基本表を出力し、ヘッダ行番号を返す。"""
    row = start_row
    if title_row:
        ws.cell(row=row, column=1, value="都道府県別人口データ")
        row += 1

    ws.cell(row=row, column=1, value="都道府県")
    for index, year in enumerate(YEARS, start=2):
        ws.cell(row=row, column=index, value=f"{year}年")

    for pref_index, prefecture in enumerate(PREFECTURES, start=row + 1):
        ws.cell(row=pref_index, column=1, value=prefecture)
        for value_index, value in enumerate(POP_DATA[pref_index - row - 1], start=2):
            ws.cell(row=pref_index, column=value_index, value=value)

    return row


def base_rows(*, english: bool = False) -> list[list[object]]:
    if english:
        headers = ["Prefecture", *[str(year) for year in YEARS]]
        prefectures = [
            "Hokkaido Prefecture",
            "Aomori Prefecture",
            "Iwate Prefecture",
            "Miyagi Prefecture",
            "Akita Prefecture",
        ]
    else:
        headers = ["都道府県", *[f"{year}年" for year in YEARS]]
        prefectures = PREFECTURES

    rows = [headers]
    for prefecture, values in zip(prefectures, POP_DATA, strict=True):
        rows.append([prefecture, *values])
    return rows


def write_rows(ws, rows: list[list[object]]) -> None:
    for row_index, row in enumerate(rows, start=1):
        for col_index, value in enumerate(row, start=1):
            ws.cell(row=row_index, column=col_index, value=value)


def set_cell_value(ws, row: int, column: int, value: object) -> None:
    cell = ws.cell(row=row, column=column, value=value)
    if isinstance(value, str) and "\n" in value:
        alignment = copy(cell.alignment)
        alignment.wrap_text = True
        cell.alignment = alignment


def save_workbook(output_dir: Path, workbook: Workbook, stem: str) -> Path:
    path = output_dir / f"{stem}.xlsx"
    workbook.save(path)
    return path


def write_delimited_file(
    output_dir: Path,
    filename: str,
    rows: list[list[object]],
    *,
    delimiter: str,
    encoding: str,
) -> Path:
    path = output_dir / filename
    with path.open("w", encoding=encoding, newline="") as file:
        writer = csv.writer(file, delimiter=delimiter)
        writer.writerows(rows)
    return path


def write_raw_text_file(output_dir: Path, filename: str, content: str, *, encoding: str = "utf-8") -> Path:
    path = output_dir / filename
    path.write_text(content, encoding=encoding)
    return path


def write_docx_table(output_dir: Path, filename: str, rows: list[list[object]]) -> Path:
    path = output_dir / filename
    document = Document()
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            table.cell(row_index, col_index).text = str(value)
    document.save(path)
    return path


def write_pdf_table(output_dir: Path, filename: str, rows: list[list[object]]) -> Path:
    path = output_dir / filename
    pdf = SimpleDocTemplate(str(path), pagesize=landscape(A4))
    table = Table(rows)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    pdf.build([table])
    return path

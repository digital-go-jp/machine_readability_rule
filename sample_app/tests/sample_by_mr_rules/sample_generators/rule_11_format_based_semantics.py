"""L1-11 用サンプル（書式のみでデータの意味を区別している例）。

OK-01: ヘッダのみ太字。データ行は既定書式で均一。
OK-02: OK-01 と同一データ。ヘッダ行は単一背景色、body 行は別の単一背景色（列内で複数色にしない）。
NG-01: 1 列でフォント色が複数（黒と赤）。
NG-02: 数値列で一部のセルのみ太字。
NG-03: 1 列で非既定の背景色が複数種類。
NG-04: 行ブロックごとに背景色を変えて区分。
NG-05: 年・売上表。売上の下位 2 行（800, 900）のみ斜体で強調。
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .common import base_rows, write_rows


def _header_font() -> Font:
    return Font(bold=True)


def _font_red() -> Font:
    return Font(color="FFFF0000")


def _fill_green() -> PatternFill:
    return PatternFill(fill_type="solid", fgColor="FF00FF00")


def _fill_red() -> PatternFill:
    return PatternFill(fill_type="solid", fgColor="FFFF0000")


def _fill_yellow() -> PatternFill:
    return PatternFill(fill_type="solid", fgColor="FFFFFF00")


def _fill_header_soft() -> PatternFill:
    """R1-11 NG-1 ヘッダ付近の薄い塗りに近い色。"""
    return PatternFill(fill_type="solid", fgColor="FFF2F2F2")


def _fill_row_band_a() -> PatternFill:
    return PatternFill(fill_type="solid", fgColor="FFDAE9F8")


def _fill_row_band_b() -> PatternFill:
    return PatternFill(fill_type="solid", fgColor="FFFFF4CC")


def _thin_border() -> Side:
    return Side(style="thin", color="FFB4B4B4")


def _write_ok01(ws) -> None:
    rows = base_rows()
    write_rows(ws, rows)
    for c in range(1, len(rows[0]) + 1):
        ws.cell(row=1, column=c).font = _header_font()


def _write_ok02(ws) -> None:
    """OK-01 と同じ表。ヘッダ・body それぞれ領域内は単一の背景色（L1-11 は body 列内の複数非既定色のみ NG）。"""
    rows = base_rows()
    write_rows(ws, rows)
    ncols = len(rows[0])
    nrows = len(rows)
    hdr_fill = _fill_header_soft()
    body_fill = _fill_row_band_a()
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = _header_font()
        cell.fill = hdr_fill
    for r in range(2, nrows + 1):
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).fill = body_fill


def _write_ng01(ws) -> None:
    """列内でフォント色が複数（L1-11 テストと同型）。"""
    rows = [
        ["ステータス", "値"],
        ["正常", 100],
        ["異常", 50],
        ["正常", 80],
    ]
    write_rows(ws, rows)
    for c in range(1, 3):
        ws.cell(row=1, column=c).font = _header_font()
    ws.cell(row=2, column=1).font = Font()  # 既定（黒）
    ws.cell(row=3, column=1).font = _font_red()
    ws.cell(row=4, column=1).font = Font()


def _write_ng02(ws) -> None:
    """数値列で 1 セルのみ太字。

    先頭データ行だけ太字にすると、その行がヘッダー行と誤検出され body 対象外になるため、
    2 行目のデータにのみ太字を付ける（LayoutDetector は連続するヘッダ候補のみ採用）。
    """
    rows = [
        ["名前", "点数"],
        ["太郎", 100],
        ["花子", 50],
        ["次郎", 80],
    ]
    write_rows(ws, rows)
    for c in range(1, 3):
        ws.cell(row=1, column=c).font = _header_font()
    ws.cell(row=3, column=2).font = Font(bold=True)


def _write_ng03(ws) -> None:
    """1 列で非既定背景が複数色。"""
    rows = [
        ["ステータス", "値"],
        ["正常", 100],
        ["異常", 50],
        ["警告", 80],
    ]
    write_rows(ws, rows)
    for c in range(1, 3):
        ws.cell(row=1, column=c).font = _header_font()
    ws.cell(row=2, column=1).fill = _fill_green()
    ws.cell(row=3, column=1).fill = _fill_red()
    ws.cell(row=4, column=1).fill = _fill_yellow()


def _write_ng04(ws) -> None:
    """行ブロック単位で背景色を切り替え（全列に帯状の塗り）。"""
    rows = [
        ["商品名", "売上(千円)"],
        ["りんご", 120],
        ["みかん", 85],
        ["大根", 72],
        ["きゅうり", 50],
        ["キャベツ", 110],
    ]
    write_rows(ws, rows)
    thin = _thin_border()
    grid = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )
    for r in range(1, 7):
        for c in range(1, 3):
            ws.cell(row=r, column=c).border = grid

    hdr = _fill_header_soft()
    band_a = _fill_row_band_a()
    band_b = _fill_row_band_b()
    for c in range(1, 3):
        ws.cell(row=1, column=c).fill = hdr
    for r in (2, 3):
        for c in range(1, 3):
            ws.cell(row=r, column=c).fill = band_a
    for r in (4, 5, 6):
        for c in range(1, 3):
            ws.cell(row=r, column=c).fill = band_b


def _write_ng05(ws) -> None:
    """年・売上表。B4・B5 の売上のみ斜体"""
    rows = [
        ["年", "売上"],
        [2024, 400],
        [2025, 650],
        [2026, 800],
        [2027, 900],
    ]
    write_rows(ws, rows)
    black = Side(style="thin", color="FF000000")
    grid = Border(left=black, right=black, top=black, bottom=black)
    for r in range(1, 6):
        for c in range(1, 3):
            ws.cell(row=r, column=c).border = grid

    for c in range(1, 3):
        ws.cell(row=1, column=c).fill = _fill_header_soft()

    align_right = Alignment(horizontal="right", vertical="center")
    for r in range(2, 6):
        for c in range(1, 3):
            ws.cell(row=r, column=c).alignment = align_right

    for r in (4, 5):
        ws.cell(row=r, column=2).font = Font(italic=True)


def generate(output_dir: Path) -> list[Path]:
    path = output_dir / "rule_11_format_based_semantics.xlsx"
    wb = Workbook()

    ws_ok = wb.active
    ws_ok.title = "OK-01"
    _write_ok01(ws_ok)

    ws_ok2 = wb.create_sheet("OK-02")
    _write_ok02(ws_ok2)

    ws_ng1 = wb.create_sheet("NG-01")
    _write_ng01(ws_ng1)

    ws_ng2 = wb.create_sheet("NG-02")
    _write_ng02(ws_ng2)

    ws_ng3 = wb.create_sheet("NG-03")
    _write_ng03(ws_ng3)

    ws_ng4 = wb.create_sheet("NG-04")
    _write_ng04(ws_ng4)

    ws_ng5 = wb.create_sheet("NG-05")
    _write_ng05(ws_ng5)

    wb.save(path)
    return [path]

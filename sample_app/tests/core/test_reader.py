"""reader.py のユニットテスト

パフォーマンス改善（read_only モード、行列上限）のテストを含む。
"""

from __future__ import annotations

import io
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

import openpyxl
import pytest

from harunobu.core.reader import (
    _SINGLE_PASS_THRESHOLD,
    DEFAULT_MAX_COLS,
    DEFAULT_MAX_ROWS,
    CorruptedFileError,
    UnsupportedFormatError,
    _convert_cell_readonly,
    _extract_sheet_metadata,
    _read_xlsx_single_pass,
    _read_xlsx_two_pass,
    read_file,
    read_workbook,
)


@pytest.fixture
def tmp_xlsx(tmp_path: Path) -> Path:
    """基本的なテスト用 .xlsx ファイルを作成する。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "ヘッダー1"
    ws["B1"] = "ヘッダー2"
    ws["A2"] = 100
    ws["B2"] = 200
    ws["A3"] = "テスト"
    ws["B3"] = None  # 空セル
    path = tmp_path / "test.xlsx"
    wb.save(str(path))
    wb.close()
    return path


@pytest.fixture
def tmp_xlsx_merged(tmp_path: Path) -> Path:
    """結合セル付きのテスト用 .xlsx ファイルを作成する。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "結合ヘッダー"
    ws.merge_cells("A1:C1")
    ws["A2"] = 1
    ws["B2"] = 2
    ws["C2"] = 3
    path = tmp_path / "merged.xlsx"
    wb.save(str(path))
    wb.close()
    return path


@pytest.fixture
def tmp_xlsx_large(tmp_path: Path) -> Path:
    """大きめのテスト用 .xlsx ファイルを作成する（50行x10列）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in range(1, 51):
        for c in range(1, 11):
            ws.cell(row=r, column=c, value=f"R{r}C{c}")
    path = tmp_path / "large.xlsx"
    wb.save(str(path))
    wb.close()
    return path


class TestReadFile:
    """read_file() の基本テスト"""

    def test_read_xlsx_basic(self, tmp_xlsx: Path):
        wb = read_file(tmp_xlsx)
        assert wb.file_format == "xlsx"
        assert wb.file_name == "test.xlsx"
        assert len(wb.sheets) == 1
        sheet = wb.sheets[0]
        assert sheet.name == "Sheet1"
        # A1 のヘッダーが読めること
        cell_a1 = sheet.get_cell(1, 1)
        assert cell_a1 is not None
        assert cell_a1.value == "ヘッダー1"
        # A2 の数値が読めること
        cell_a2 = sheet.get_cell(2, 1)
        assert cell_a2 is not None
        assert cell_a2.value == 100

    def test_read_xlsx_merged_cells(self, tmp_xlsx_merged: Path):
        wb = read_file(tmp_xlsx_merged)
        sheet = wb.sheets[0]
        # 結合セル情報が取得できること
        assert len(sheet.merged_cells) == 1
        mr = sheet.merged_cells[0]
        assert mr.start_row == 1
        assert mr.start_col == 1
        assert mr.end_row == 1
        assert mr.end_col == 3
        # 結合セルのマスター情報
        cell_a1 = sheet.get_cell(1, 1)
        assert cell_a1 is not None
        assert cell_a1.is_merged is True
        assert cell_a1.value == "結合ヘッダー"

    def test_read_xlsx_skips_empty_cells(self, tmp_xlsx: Path):
        wb = read_file(tmp_xlsx)
        sheet = wb.sheets[0]
        # B3 は None なのでスキップされるべき
        cell_b3 = sheet.get_cell(3, 2)
        assert cell_b3 is None

    def test_file_not_found(self):
        with pytest.raises(FileNotFoundError):
            read_file("/nonexistent/file.xlsx")

    def test_unsupported_format_xls(self, tmp_path: Path):
        p = tmp_path / "old.xls"
        p.write_bytes(b"dummy")
        with pytest.raises(UnsupportedFormatError):
            read_file(p)

    def test_unsupported_format_unknown(self, tmp_path: Path):
        p = tmp_path / "file.dat"
        p.write_bytes(b"dummy")
        with pytest.raises(UnsupportedFormatError):
            read_file(p)


class TestReadWorkbookContentValidation:
    """read_workbook() の bytes 入力検証テスト"""

    def test_read_workbook_rejects_xls(self):
        with pytest.raises(UnsupportedFormatError, match=r"\.xls 形式はサポートされていません"):
            read_workbook("old.xls", io.BytesIO(b"dummy"))

    def test_read_workbook_rejects_spoofed_xlsx_bytes(self, monkeypatch):
        class FakeMagika:
            def identify_bytes(self, content: bytes):
                assert content == b"%PDF-1.7\n"
                return SimpleNamespace(output=SimpleNamespace(label="pdf"))

        monkeypatch.setattr("harunobu.core.reader._magika_instance", FakeMagika())

        with pytest.raises(UnsupportedFormatError, match="拡張子が偽装"):
            read_workbook("spoofed.xlsx", io.BytesIO(b"%PDF-1.7\n"))

    def test_read_workbook_corrupted_xlsx_raises_user_facing_error(self, monkeypatch):
        class FakeMagika:
            def identify_bytes(self, _content: bytes):
                return SimpleNamespace(output=SimpleNamespace(label="zip"))

        monkeypatch.setattr("harunobu.core.reader._magika_instance", FakeMagika())

        with pytest.raises(CorruptedFileError, match="破損の可能性"):
            read_workbook("broken.xlsx", io.BytesIO(b"not a valid workbook"))

    def test_read_workbook_falls_back_when_magika_bytes_detection_fails(self, monkeypatch):
        class FailingMagika:
            def identify_bytes(self, _content: bytes):
                raise RuntimeError("magika failed")

        monkeypatch.setattr("harunobu.core.reader._magika_instance", FailingMagika())

        result = read_workbook("fallback.csv", io.BytesIO(b"a,b\n1,2\n"))

        assert result.file_format == "csv"
        assert result.sheets[0].get_cell_value(2, 2) == "2"


class TestFormulaArtifacts:
    """OOXML 数式メタデータの統合テスト。"""

    def test_read_file_attaches_formula_refs(self, tmp_path: Path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "値"
        ws["A2"] = 10
        ws["B2"] = "=A2"
        path = tmp_path / "formula.xlsx"
        wb.save(str(path))
        wb.close()

        result = read_file(path)
        sheet = result.sheets[0]

        assert sheet.formula_refs_loaded is True
        assert sheet.formula_refs_error is None
        assert [(ref.cell_ref, ref.formula_text) for ref in sheet.formula_refs] == [("B2", "A2")]

    def test_read_file_xlsm_attaches_formula_refs(self, tmp_path: Path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "MacroLike"
        ws["A1"] = 1
        ws["B1"] = "=A1"
        path = tmp_path / "formula.xlsm"
        wb.save(str(path))
        wb.close()

        result = read_file(path)

        assert result.sheets[0].formula_refs_loaded is True
        assert result.sheets[0].formula_refs[0].cell_ref == "B1"

    def test_read_workbook_bytes_attaches_formula_refs(self, tmp_path: Path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Bytes"
        ws["A1"] = 1
        ws["B1"] = "=A1"
        path = tmp_path / "bytes.xlsx"
        wb.save(str(path))
        wb.close()

        result = read_workbook("bytes.xlsx", io.BytesIO(path.read_bytes()))

        assert result.sheets[0].formula_refs_loaded is True
        assert result.sheets[0].formula_refs[0].cell_ref == "B1"

    def test_no_formula_file_marks_formula_refs_loaded(self, tmp_xlsx: Path):
        result = read_file(tmp_xlsx)

        assert result.sheets[0].formula_refs_loaded is True
        assert result.sheets[0].formula_refs == []
        assert result.sheets[0].formula_error_refs == []

    def test_inspector_error_does_not_abort_reader(self, tmp_xlsx: Path, monkeypatch):
        class FailedResult:
            formula_refs = []
            formula_error_refs = []
            error = "too_many_formula_artifacts"

        monkeypatch.setattr(
            "harunobu.core.reader.ooxml_inspector.inspect_formula_artifacts",
            lambda _source: FailedResult(),
        )

        result = read_file(tmp_xlsx)

        assert result.sheets[0].formula_refs_loaded is False
        assert result.sheets[0].formula_refs_error == "too_many_formula_artifacts"


class TestRowColumnLimits:
    """行・列上限パラメータのテスト"""

    def test_default_limits(self):
        assert DEFAULT_MAX_ROWS == 10_000
        assert DEFAULT_MAX_COLS == 500

    def test_max_rows_limit(self, tmp_xlsx_large: Path):
        # 10行に制限して読み込み
        wb = read_file(tmp_xlsx_large, max_rows=10)
        sheet = wb.sheets[0]
        # 10行目のデータは存在するが、11行目以降は読み込まれない
        cell_r10 = sheet.get_cell(10, 1)
        assert cell_r10 is not None
        assert cell_r10.value == "R10C1"
        # 11行目のセルは読み込まれていない（結合セル情報がなければ）
        cell_r11 = sheet.get_cell(11, 1)
        assert cell_r11 is None

    def test_max_cols_limit(self, tmp_xlsx_large: Path):
        # 5列に制限して読み込み
        wb = read_file(tmp_xlsx_large, max_cols=5)
        sheet = wb.sheets[0]
        cell_c5 = sheet.get_cell(1, 5)
        assert cell_c5 is not None
        assert cell_c5.value == "R1C5"
        # 6列目以降は読み込まれない
        cell_c6 = sheet.get_cell(1, 6)
        assert cell_c6 is None

    def test_custom_limits(self, tmp_xlsx_large: Path):
        wb = read_file(tmp_xlsx_large, max_rows=5, max_cols=3)
        sheet = wb.sheets[0]
        # 5行3列以内のデータのみ存在
        assert sheet.get_cell(5, 3) is not None
        assert sheet.get_cell(5, 3).value == "R5C3"
        assert sheet.get_cell(6, 1) is None
        assert sheet.get_cell(1, 4) is None


class TestExtractSheetMetadata:
    """_extract_sheet_metadata のテスト"""

    def test_metadata_includes_merged_cells(self, tmp_xlsx_merged: Path):
        wb = openpyxl.load_workbook(str(tmp_xlsx_merged), data_only=True)
        ws = wb.worksheets[0]
        meta = _extract_sheet_metadata(ws)
        wb.close()

        assert len(meta["merged_cells"]) == 1
        assert meta["merged_cells"][0].start_col == 1
        assert meta["merged_cells"][0].end_col == 3

    def test_metadata_merge_masters(self, tmp_xlsx_merged: Path):
        wb = openpyxl.load_workbook(str(tmp_xlsx_merged), data_only=True)
        ws = wb.worksheets[0]
        meta = _extract_sheet_metadata(ws)
        wb.close()

        # A1:C1 の結合 → (1,1), (1,2), (1,3) がマスター (1,1) を指す
        masters = meta["merge_masters"]
        assert (1, 1) in masters
        assert (1, 2) in masters
        assert (1, 3) in masters
        assert masters[(1, 2)].row == 1
        assert masters[(1, 2)].col == 1


class TestConvertCellReadonly:
    """_convert_cell_readonly のテスト"""

    def test_basic_value(self):
        cell = MagicMock()
        cell.row = 1
        cell.column = 2
        cell.value = "hello"
        cell.number_format = "General"

        result = _convert_cell_readonly(cell)
        assert result is not None
        assert result.pos.row == 1
        assert result.pos.col == 2
        assert result.value == "hello"

    def test_formula_detection(self):
        cell = MagicMock()
        cell.row = 1
        cell.column = 1
        cell.value = "=SUM(A1:A10)"
        cell.number_format = "General"

        result = _convert_cell_readonly(cell)
        assert result is not None
        assert result.formula == "=SUM(A1:A10)"
        assert result.value is None

    def test_number_format(self):
        cell = MagicMock()
        cell.row = 1
        cell.column = 1
        cell.value = 42
        cell.number_format = "#,##0"

        result = _convert_cell_readonly(cell)
        assert result is not None
        assert result.fmt.number_format == "#,##0"


class TestCSVReading:
    """CSV/TSV 読み込みのテスト"""

    def test_csv_still_works(self, tmp_path: Path):
        csv_path = tmp_path / "test.csv"
        csv_path.write_text("a,b,c\n1,2,3\n", encoding="utf-8")
        wb = read_file(csv_path)
        assert wb.file_format == "csv"
        assert wb.sheets[0].get_cell_value(1, 1) == "a"
        assert wb.sheets[0].get_cell_value(2, 2) == "2"

    def test_tsv_still_works(self, tmp_path: Path):
        tsv_path = tmp_path / "test.tsv"
        tsv_path.write_text("x\ty\n10\t20\n", encoding="utf-8")
        wb = read_file(tsv_path)
        assert wb.file_format == "tsv"
        assert wb.sheets[0].get_cell_value(1, 1) == "x"

    def test_csv_bom_stripped(self, tmp_path: Path):
        """BOM付きUTF-8 CSVが正しく読み込まれること（BOMが値に残らない）"""
        csv_path = tmp_path / "bom.csv"
        csv_path.write_bytes(b"\xef\xbb\xbf" + '"名前","住所"\n"太郎","東京"\n'.encode("utf-8"))
        wb = read_file(csv_path)
        assert wb.file_format == "csv"
        # BOM がフィールド値に混入しないこと
        first_val = wb.sheets[0].get_cell_value(1, 1)
        assert first_val == "名前", f"BOM残留: {repr(first_val)}"
        assert "\ufeff" not in str(first_val)

    def test_csv_bom_quoting_preserved(self, tmp_path: Path):
        """BOM付きCSVでクォーティングが正しく解析されること"""
        csv_path = tmp_path / "bom_quoted.csv"
        # BOM + クォートされたフィールド
        csv_path.write_bytes(b"\xef\xbb\xbf" + '"都道府県","人口"\n"東京都","14,000,000"\n'.encode("utf-8"))
        wb = read_file(csv_path)
        # クォーティングが正しく処理され、カンマを含む値がそのまま読めること
        assert wb.sheets[0].get_cell_value(1, 1) == "都道府県"
        assert wb.sheets[0].get_cell_value(2, 2) == "14,000,000"

    def test_csv_cp932_fallback(self, tmp_path: Path):
        """cp932 エンコーディングのCSVが読み込めること"""
        csv_path = tmp_path / "sjis.csv"
        csv_path.write_bytes("名前,住所\n太郎,東京\n".encode("cp932"))
        wb = read_file(csv_path)
        assert wb.file_format == "csv"
        assert wb.sheets[0].get_cell_value(1, 1) == "名前"

    def test_csv_empty_fields_are_none(self, tmp_path: Path):
        """空フィールドがNoneに変換されること"""
        csv_path = tmp_path / "empty.csv"
        csv_path.write_text("a,,c\n1,,3\n", encoding="utf-8")
        wb = read_file(csv_path)
        assert wb.sheets[0].get_cell_value(1, 2) is None
        assert wb.sheets[0].get_cell_value(2, 2) is None

    def test_csv_diagnostics_accept_excel_style_selective_quotes(self, tmp_path: Path):
        """Excel由来の必要箇所だけクォートされたCSVは問題にしない"""
        csv_path = tmp_path / "selective_quotes.csv"
        csv_path.write_text('id,name,value\n1,"Tokyo, Japan",100\n2,Osaka,200\n', encoding="utf-8")
        wb = read_file(csv_path)
        diagnostics = wb.sheets[0].csv_diagnostics
        assert diagnostics is not None
        assert diagnostics.expected_field_count == 3
        assert diagnostics.issue_counts == {}
        assert diagnostics.truncated is False
        assert wb.sheets[0].get_cell_value(2, 2) == "Tokyo, Japan"

    def test_csv_diagnostics_detects_last_row_unquoted_comma(self, tmp_path: Path):
        """最後の行だけ未クォートカンマで列数が崩れる境界を検出する"""
        csv_path = tmp_path / "last_row_bad.csv"
        rows = ["id,name,value"]
        rows.extend(f"{i},Name {i},{i}" for i in range(1, 20))
        rows.append("20,Tokyo, Japan,100")
        csv_path.write_text("\n".join(rows) + "\n", encoding="utf-8")
        wb = read_file(csv_path)
        diagnostics = wb.sheets[0].csv_diagnostics
        assert diagnostics is not None
        assert diagnostics.truncated is False
        assert diagnostics.issue_counts["field_count_mismatch"] == 1
        sample = diagnostics.issue_samples[0]
        assert sample.issue_type == "field_count_mismatch"
        assert sample.logical_row == 21
        assert sample.field_count == 4
        assert sample.expected_field_count == 3

    def test_csv_diagnostics_detects_unquoted_double_quote(self, tmp_path: Path):
        """英文中の未クォートdouble quoteをraw診断で検出する"""
        csv_path = tmp_path / "unquoted_quote.csv"
        csv_path.write_text('id,comment\n1,He said "OK"\n', encoding="utf-8")
        wb = read_file(csv_path)
        diagnostics = wb.sheets[0].csv_diagnostics
        assert diagnostics is not None
        assert diagnostics.issue_counts["unquoted_quote"] == 1
        assert diagnostics.issue_samples[0].cell_range == "B2"

    def test_csv_diagnostics_detects_unclosed_quote(self, tmp_path: Path):
        """quote閉じ忘れはparse errorとして検出する"""
        csv_path = tmp_path / "unclosed.csv"
        csv_path.write_text('id,comment\n1,"unclosed\n2,next\n', encoding="utf-8")
        wb = read_file(csv_path)
        diagnostics = wb.sheets[0].csv_diagnostics
        assert diagnostics is not None
        assert diagnostics.issue_counts["unclosed_quote"] == 1
        sample = diagnostics.issue_samples[0]
        assert sample.logical_row == 2
        assert sample.physical_start_line == 2
        assert sample.physical_end_line == 3

    def test_csv_diagnostics_bom_and_cp932_routes_create_diagnostics(self, tmp_path: Path):
        """BOM付きUTF-8とcp932の既存ルーティングでも診断が作られる"""
        bom_path = tmp_path / "bom_diag.csv"
        bom_path.write_bytes(b"\xef\xbb\xbf" + '"名前","住所"\n"太郎","東京, 日本"\n'.encode("utf-8"))
        bom_wb = read_file(bom_path)
        bom_diag = bom_wb.sheets[0].csv_diagnostics
        assert bom_diag is not None
        assert bom_diag.expected_field_count == 2
        assert bom_diag.issue_counts == {}

        cp932_path = tmp_path / "cp932_diag.csv"
        cp932_path.write_bytes('名前,備考\n太郎,He said "OK"\n'.encode("cp932"))
        cp932_wb = read_file(cp932_path)
        cp932_diag = cp932_wb.sheets[0].csv_diagnostics
        assert cp932_diag is not None
        assert cp932_diag.issue_counts["unquoted_quote"] == 1

    def test_csv_diagnostics_truncates_after_many_issues(self, tmp_path: Path, monkeypatch):
        """大量違反では詳細保持を打ち切り、通常セル読込は続ける"""
        monkeypatch.setattr("harunobu.core.reader.MAX_CSV_ISSUES_BEFORE_ABORT", 3)
        monkeypatch.setattr("harunobu.core.reader.MAX_CSV_ISSUE_SAMPLES", 2)
        csv_path = tmp_path / "many_bad_rows.csv"
        csv_path.write_text(
            'id,comment\n1,He said "A"\n2,He said "B"\n3,He said "C"\n4,He said "D"\n5,still read\n',
            encoding="utf-8",
        )
        wb = read_file(csv_path)
        diagnostics = wb.sheets[0].csv_diagnostics
        assert diagnostics is not None
        assert diagnostics.truncated is True
        assert diagnostics.abort_reason == "too_many_issues"
        assert diagnostics.issue_counts["unquoted_quote"] == 4
        assert len(diagnostics.issue_samples) == 2
        assert wb.sheets[0].get_cell_value(6, 2) == "still read"


class TestTwoPassApproach:
    """2パス方式（メタデータ+read_only）の統合テスト"""

    def test_hidden_rows_preserved(self, tmp_path: Path):
        """非表示行の情報がメタデータパスで取得できること"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "visible"
        ws["A2"] = "hidden"
        ws["A3"] = "visible"
        ws.row_dimensions[2].hidden = True
        path = tmp_path / "hidden.xlsx"
        wb.save(str(path))
        wb.close()

        result = read_file(path)
        sheet = result.sheets[0]
        assert 2 in sheet.hidden_rows

    def test_hidden_cols_preserved(self, tmp_path: Path):
        """非表示列の情報がメタデータパスで取得できること"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "a"
        ws["B1"] = "b"
        ws.column_dimensions["B"].hidden = True
        path = tmp_path / "hidden_col.xlsx"
        wb.save(str(path))
        wb.close()

        result = read_file(path)
        sheet = result.sheets[0]
        assert 2 in sheet.hidden_cols

    def test_single_pass_and_two_pass_produce_same_result(self, tmp_path: Path):
        """1パス方式と2パス方式で同一のセルデータが得られること"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        ws["A1"] = "名前"
        ws["B1"] = "値"
        ws["A2"] = "項目A"
        ws["B2"] = 12345
        ws["A3"] = "項目B"
        ws["B3"] = None
        ws.merge_cells("A4:B4")
        ws["A4"] = "結合セル"
        ws.row_dimensions[3].hidden = True
        ws.column_dimensions["B"].hidden = True
        path = tmp_path / "compare.xlsx"
        wb.save(str(path))
        wb.close()

        data = path.read_bytes()
        max_rows, max_cols = DEFAULT_MAX_ROWS, DEFAULT_MAX_COLS

        result_single = _read_xlsx_single_pass(data, "compare.xlsx", max_rows=max_rows, max_cols=max_cols)
        result_two = _read_xlsx_two_pass(data, "compare.xlsx", max_rows=max_rows, max_cols=max_cols)

        assert len(result_single.sheets) == len(result_two.sheets)
        s1, s2 = result_single.sheets[0], result_two.sheets[0]
        assert s1.name == s2.name
        # セルデータの一致
        for r in range(1, 5):
            for c in range(1, 3):
                assert s1.get_cell_value(r, c) == s2.get_cell_value(r, c), (
                    f"Cell ({r},{c}) differs: {s1.get_cell_value(r, c)} vs {s2.get_cell_value(r, c)}"
                )
        # メタデータの一致
        assert s1.hidden_rows == s2.hidden_rows
        assert s1.hidden_cols == s2.hidden_cols
        assert s1.merged_cells == s2.merged_cells

    def test_threshold_routes_correctly(self, tmp_path: Path, monkeypatch):
        """ファイルサイズに応じて1パス/2パスが切り替わること"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "test"
        path = tmp_path / "small.xlsx"
        wb.save(str(path))
        wb.close()

        # 小さいファイル（閾値以下）→ 1パス
        assert path.stat().st_size <= _SINGLE_PASS_THRESHOLD
        result = read_file(path)
        assert result.sheets[0].get_cell_value(1, 1) == "test"

        # 閾値を0にして強制的に2パスを通す
        monkeypatch.setattr("harunobu.core.reader._SINGLE_PASS_THRESHOLD", 0)
        result_two = read_file(path)
        assert result_two.sheets[0].get_cell_value(1, 1) == "test"


class TestAttachObjectArtifacts:
    """_attach_object_artifacts が Sheet.object_refs / object_refs_loaded を正しく設定する。"""

    def _make_xlsx_bytes_no_drawing(self) -> bytes:
        """drawing なしの最小 xlsx バイト列を返す。"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "test"
        buf = io.BytesIO()
        wb.save(buf)
        wb.close()
        return buf.getvalue()

    def test_object_refs_loaded_true_on_success(self):
        """drawing なし xlsx で object_refs_loaded=True になる。"""
        data = self._make_xlsx_bytes_no_drawing()
        result = _read_xlsx_single_pass(data, "test.xlsx", max_rows=100, max_cols=100)
        sheet = result.sheets[0]
        assert sheet.object_refs_loaded is True
        assert sheet.object_refs_error is None
        assert sheet.object_refs == []

    def test_object_refs_loaded_via_two_pass(self):
        """2パス方式でも object_refs_loaded が設定される。"""
        data = self._make_xlsx_bytes_no_drawing()
        result = _read_xlsx_two_pass(data, "test.xlsx", max_rows=100, max_cols=100)
        sheet = result.sheets[0]
        assert sheet.object_refs_loaded is True
        assert sheet.object_refs_error is None

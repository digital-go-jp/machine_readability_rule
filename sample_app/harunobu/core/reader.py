"""Excel / CSV / TSV ファイルリーダー

ファイルを読み込み、内部データモデル (WorkBook) に変換する。
"""

from __future__ import annotations

import csv
import io
import logging
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import IO, Any

import openpyxl
from openpyxl.cell.cell import Cell as OpenpyxlCell
from openpyxl.worksheet.worksheet import Worksheet

from harunobu.core import ooxml_inspector
from harunobu.core.models import (
    Cell,
    CellFormat,
    CellPosition,
    CsvDiagnostics,
    CsvIssueSample,
    MergedRange,
    ObjectRef,
    Sheet,
    WorkBook,
)

logger = logging.getLogger(__name__)

# ─── Magika によるファイル実体検証 ───

try:
    from magika import Magika as _Magika

    _magika_instance: _Magika | None = _Magika()
except Exception:
    _magika_instance = None

# 拡張子 → 許可される Magika ラベルのマッピング
_EXTENSION_ALLOWED_LABELS: dict[str, set[str]] = {
    ".xlsx": {"xlsx", "zip"},
    ".xlsm": {"xlsx", "zip"},
    ".csv": {"csv", "txt"},
    ".tsv": {"csv", "tsv", "txt"},
}


def detect_content_label(path: Path) -> str | None:
    """Magika でファイルの内容ラベルを判定する。

    Magika が利用できない場合は None を返す。
    """
    if _magika_instance is None:
        return None
    try:
        result = _magika_instance.identify_path(path)
        return result.output.label
    except Exception:
        logger.debug("Magika による内容判定に失敗: %s", path, exc_info=True)
        return None


def detect_content_label_bytes(content: bytes, source_name: str = "<bytes>") -> str | None:
    """Magika でバイト列の内容ラベルを判定する。

    Magika が利用できない場合は None を返す。
    """
    if _magika_instance is None:
        return None
    try:
        result = _magika_instance.identify_bytes(content)
        return result.output.label
    except Exception:
        logger.debug("Magika による内容判定に失敗: %s", source_name, exc_info=True)
        return None


def _validate_content_label(file_name: str, suffix: str, label: str | None) -> None:
    """拡張子と Magika ラベルの対応を検証する。"""
    if label is None:
        return

    allowed_labels = _EXTENSION_ALLOWED_LABELS.get(suffix)
    if allowed_labels is None:
        return

    if label not in allowed_labels:
        raise UnsupportedFormatError(
            f"ファイル '{file_name}' の拡張子は '{suffix}' ですが、"
            f"実際のファイル内容は '{label}' として検出されました。"
            f"拡張子が偽装されている可能性があります。"
        )


def validate_file_content(path: Path) -> None:
    """ファイルの実体が拡張子と一致するか検証する。

    Magika が利用できない環境では何もしない（拡張子ベースのチェックにフォールバック）。
    不一致の場合は UnsupportedFormatError を送出する。
    """
    suffix = path.suffix.lower()
    _validate_content_label(path.name, suffix, detect_content_label(path))


def validate_file_content_bytes(file_name: str, content: bytes) -> None:
    """バイト列の実体がファイル名の拡張子と一致するか検証する。

    Magika が利用できない環境では何もしない（拡張子ベースのチェックにフォールバック）。
    不一致の場合は UnsupportedFormatError を送出する。
    """
    suffix = Path(file_name).suffix.lower()
    _validate_content_label(file_name, suffix, detect_content_label_bytes(content, file_name))


# デフォルトの行・列上限（パフォーマンス保護）
DEFAULT_MAX_ROWS = 10_000
DEFAULT_MAX_COLS = 500
CSV_RAW_PREVIEW_LIMIT = 500
MAX_CSV_ISSUE_SAMPLES = 100
MAX_CSV_ISSUE_SAMPLES_PER_TYPE = 50
MAX_CSV_ISSUES_BEFORE_ABORT = 1_000


class UnsupportedFormatError(Exception):
    """サポートされていないファイル形式"""


def read_workbook(
    file_name: str,
    file_bytes: IO[bytes],
    *,
    max_rows: int = DEFAULT_MAX_ROWS,
    max_cols: int = DEFAULT_MAX_COLS,
) -> WorkBook:
    """ファイル名とバイトストリームから WorkBook を読み込む。

    Args:
        file_name (str): ファイル名（拡張子でフォーマット判定）
        file_bytes (IO[bytes]): ファイルのバイトストリーム
        max_rows (int): 読み込む最大行数
        max_cols (int): 読み込む最大列数

    Returns:
        WorkBook
    """
    suffix = Path(file_name).suffix.lower()

    if suffix == ".xls":
        raise UnsupportedFormatError(f".xls 形式はサポートされていません。.xlsx に変換してください: {file_name}")

    data = file_bytes.read()
    validate_file_content_bytes(file_name, data)

    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx_data(file_name, data, max_rows=max_rows, max_cols=max_cols)
    elif suffix == ".csv":
        return _read_csv_data(file_name, data, delimiter=",")
    elif suffix == ".tsv":
        return _read_csv_data(file_name, data, delimiter="\t")
    else:
        raise UnsupportedFormatError(f"未対応のファイル形式です: {suffix}")


def read_file(
    path: str | Path,
    *,
    max_rows: int = DEFAULT_MAX_ROWS,
    max_cols: int = DEFAULT_MAX_COLS,
) -> WorkBook:
    """ファイルを読み込み WorkBook を返す。

    Args:
        path (str | Path): 読み込むファイルパス (.xlsx, .csv, .tsv)
        max_rows (int): 読み込む最大行数（パフォーマンス保護、デフォルト 10,000）
        max_cols (int): 読み込む最大列数（パフォーマンス保護、デフォルト 500）

    Returns:
        WorkBook

    Raises:
        UnsupportedFormatError: .xls など未対応形式の場合
        FileNotFoundError: ファイルが存在しない場合
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"ファイルが見つかりません: {path}")

    suffix = path.suffix.lower()

    if suffix == ".xls":
        raise UnsupportedFormatError(f".xls 形式はサポートされていません。.xlsx に変換してください: {path.name}")

    # Magika が利用可能なら拡張子偽装を検出（パース前に弾く）
    validate_file_content(path)

    if suffix in (".xlsx", ".xlsm"):
        wb = _read_xlsx(path, max_rows=max_rows, max_cols=max_cols)
    elif suffix == ".csv":
        wb = _read_csv(path, delimiter=",")
    elif suffix == ".tsv":
        wb = _read_csv(path, delimiter="\t")
    else:
        raise UnsupportedFormatError(f"未対応のファイル形式です: {suffix}")

    wb.file_path = str(path.resolve())
    return wb


# ─── Excel (.xlsx) ───


class CorruptedFileError(Exception):
    """ファイルが破損している"""


# 5MB以下のファイルは1パス方式（通常モード）で高速に処理
_SINGLE_PASS_THRESHOLD = 5 * 1024 * 1024


def _read_xlsx_bytes(file_name: str, file_bytes: IO[bytes], *, max_rows: int, max_cols: int) -> WorkBook:
    """バイトストリームから .xlsx を読み込む。

    小さいファイル（5MB以下）は1パス方式、大きいファイルは2パス方式で処理する。
    """
    return _read_xlsx_data(file_name, file_bytes.read(), max_rows=max_rows, max_cols=max_cols)


def _read_xlsx_data(file_name: str, data: bytes, *, max_rows: int, max_cols: int) -> WorkBook:
    """バイト列から .xlsx を読み込む。"""
    if len(data) <= _SINGLE_PASS_THRESHOLD:
        return _read_xlsx_single_pass(data, file_name, max_rows=max_rows, max_cols=max_cols)
    return _read_xlsx_two_pass(data, file_name, max_rows=max_rows, max_cols=max_cols)


def _read_xlsx_single_pass(data: bytes, file_name: str, *, max_rows: int, max_cols: int) -> WorkBook:
    """通常モード1回で全情報を取得する（小ファイル向け）。"""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:
        raise CorruptedFileError(f"ファイルを開けません（破損の可能性）: {file_name}: {e}") from e

    sheets: list[Sheet] = []
    for ws in wb.worksheets:
        sheet = _parse_worksheet(ws, max_rows=max_rows, max_cols=max_cols)
        sheets.append(sheet)

    wb.close()
    workbook = WorkBook(
        file_name=file_name,
        file_format="xlsx",
        sheets=sheets,
        file_size=len(data),
    )
    _attach_formula_artifacts(workbook.sheets, data)
    _attach_object_artifacts(workbook.sheets, data)
    return workbook


def _read_xlsx_two_pass(data: bytes, file_name: str, *, max_rows: int, max_cols: int) -> WorkBook:
    """2パス方式: メタデータ（通常モード）+ セルデータ（read_onlyモード）。大ファイル向け。"""
    try:
        wb_meta = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:
        raise CorruptedFileError(f"ファイルを開けません（破損の可能性）: {file_name}: {e}") from e

    sheet_metadata: dict[str, dict] = {}
    for ws in wb_meta.worksheets:
        sheet_metadata[ws.title] = _extract_sheet_metadata(ws)
    wb_meta.close()

    try:
        wb_ro = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as e:
        raise CorruptedFileError(f"ファイルを開けません（破損の可能性）: {file_name}: {e}") from e

    sheets: list[Sheet] = []
    for ws_ro in wb_ro.worksheets:
        meta = sheet_metadata.get(ws_ro.title, {})
        sheet = _parse_worksheet_fast(ws_ro, meta, max_rows=max_rows, max_cols=max_cols)
        sheets.append(sheet)

    wb_ro.close()
    workbook = WorkBook(
        file_name=file_name,
        file_format="xlsx",
        sheets=sheets,
        file_size=len(data),
    )
    _attach_formula_artifacts(workbook.sheets, data)
    _attach_object_artifacts(workbook.sheets, data)
    return workbook


def _read_xlsx(path: Path, *, max_rows: int, max_cols: int) -> WorkBook:
    """openpyxl で .xlsx を読み込む。

    小さいファイル（5MB以下）は1パス方式、大きいファイルは2パス方式で処理する。
    """
    file_size = path.stat().st_size
    if file_size <= _SINGLE_PASS_THRESHOLD:
        # 1パス方式: 通常モードで全情報を一度に取得
        try:
            wb = openpyxl.load_workbook(str(path), data_only=True)
        except Exception as e:
            raise CorruptedFileError(f"ファイルを開けません（破損の可能性）: {path.name}: {e}") from e

        sheets: list[Sheet] = []
        for ws in wb.worksheets:
            sheet = _parse_worksheet(ws, max_rows=max_rows, max_cols=max_cols)
            sheets.append(sheet)

        workbook = WorkBook(
            file_name=path.name,
            file_format="xlsx",
            sheets=sheets,
            file_size=file_size,
        )
        wb.close()
        _attach_formula_artifacts(workbook.sheets, path)
        _attach_object_artifacts(workbook.sheets, path)
        return workbook

    # 2パス方式: 大ファイルはread_onlyモードでセルデータを効率的に読む
    try:
        wb_meta = openpyxl.load_workbook(str(path), data_only=True)
    except Exception as e:
        raise CorruptedFileError(f"ファイルを開けません（破損の可能性）: {path.name}: {e}") from e

    sheet_metadata: dict[str, dict] = {}
    for ws in wb_meta.worksheets:
        sheet_metadata[ws.title] = _extract_sheet_metadata(ws)
    wb_meta.close()

    try:
        wb_ro = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    except Exception as e:
        raise CorruptedFileError(f"ファイルを開けません（破損の可能性）: {path.name}: {e}") from e

    sheets_list: list[Sheet] = []
    for ws_ro in wb_ro.worksheets:
        meta = sheet_metadata.get(ws_ro.title, {})
        sheet = _parse_worksheet_fast(ws_ro, meta, max_rows=max_rows, max_cols=max_cols)
        sheets_list.append(sheet)

    workbook = WorkBook(
        file_name=path.name,
        file_format="xlsx",
        sheets=sheets_list,
        file_size=file_size,
    )
    wb_ro.close()
    _attach_formula_artifacts(workbook.sheets, path)
    _attach_object_artifacts(workbook.sheets, path)
    return workbook


def _attach_object_artifacts(sheets: list[Sheet], source: bytes | Path) -> None:
    """解析済みの各シートに、軽量な OOXML オブジェクトメタデータを付与する。"""
    result = ooxml_inspector.inspect_object_artifacts(source)
    if result.error is not None:
        logger.warning("OOXML オブジェクトメタデータの取得に失敗: %s", result.error)
        for sheet in sheets:
            sheet.object_refs = []
            sheet.object_refs_loaded = False
            sheet.object_refs_error = result.error
        return

    by_sheet: dict[str, list[ObjectRef]] = {s.name: [] for s in sheets}
    for ref in result.object_refs:
        by_sheet.setdefault(ref.sheet_name, []).append(ref)

    for sheet in sheets:
        sheet.object_refs = by_sheet.get(sheet.name, [])
        sheet.object_refs_loaded = True
        sheet.object_refs_error = None


def _attach_formula_artifacts(sheets: list[Sheet], source: bytes | Path) -> None:
    """解析済みの各シートに、軽量な OOXML 数式メタデータを付与する。"""
    result = ooxml_inspector.inspect_formula_artifacts(source)
    if result.error is not None:
        logger.warning("OOXML 数式メタデータの取得に失敗: %s", result.error)
        for sheet in sheets:
            sheet.formula_refs = []
            sheet.formula_error_refs = []
            sheet.formula_refs_loaded = False
            sheet.formula_refs_error = result.error
        return

    formula_refs_by_sheet = {sheet.name: [] for sheet in sheets}
    formula_error_refs_by_sheet = {sheet.name: [] for sheet in sheets}

    for formula_ref in result.formula_refs:
        formula_refs_by_sheet.setdefault(formula_ref.sheet_name, []).append(formula_ref)
    for error_ref in result.formula_error_refs:
        formula_error_refs_by_sheet.setdefault(error_ref.sheet_name, []).append(error_ref)

    for sheet in sheets:
        sheet.formula_refs = formula_refs_by_sheet.get(sheet.name, [])
        sheet.formula_error_refs = formula_error_refs_by_sheet.get(sheet.name, [])
        sheet.formula_refs_loaded = True
        sheet.formula_refs_error = None


def _extract_sheet_metadata(ws: Worksheet) -> dict:
    """通常モードの Worksheet からメタデータ（結合セル、非表示行列、グループ、オブジェクト等）を取得する。"""
    # 結合セル
    merged_cells: list[MergedRange] = []
    merge_masters: dict[tuple[int, int], CellPosition] = {}

    for merged_range in ws.merged_cells.ranges:
        mr = MergedRange(
            start_row=merged_range.min_row,
            start_col=merged_range.min_col,
            end_row=merged_range.max_row,
            end_col=merged_range.max_col,
        )
        merged_cells.append(mr)

        master_pos = CellPosition(row=mr.start_row, col=mr.start_col)
        for r in range(mr.start_row, mr.end_row + 1):
            for c in range(mr.start_col, mr.end_col + 1):
                merge_masters[(r, c)] = master_pos

    # 非表示行
    hidden_rows: list[int] = []
    if ws.row_dimensions:
        for row_idx, rd in ws.row_dimensions.items():
            if rd.hidden:
                hidden_rows.append(int(row_idx))

    # 非表示列
    hidden_cols: list[int] = []
    if ws.column_dimensions:
        for col_key, cd in ws.column_dimensions.items():
            if cd.hidden:
                if isinstance(col_key, str):
                    col_idx = openpyxl.utils.column_index_from_string(col_key)
                else:
                    col_idx = int(col_key)
                hidden_cols.append(col_idx)

    # 行グループ / 列グループ
    row_groups = _extract_row_groups(ws)
    col_groups = _extract_col_groups(ws)

    return {
        "merged_cells": merged_cells,
        "merge_masters": merge_masters,
        "hidden_rows": sorted(hidden_rows),
        "hidden_cols": sorted(hidden_cols),
        "row_groups": row_groups,
        "col_groups": col_groups,
        "sheet_state": ws.sheet_state,
        "max_row": ws.max_row or 0,
        "max_col": ws.max_column or 0,
    }


def _parse_worksheet_fast(ws, meta: dict, *, max_rows: int, max_cols: int) -> Sheet:
    """read_only モードの Worksheet とメタデータから Sheet モデルを構築する。"""
    cells: dict[tuple[int, int], Cell] = {}

    meta_max_row = meta.get("max_row", 0)
    meta_max_col = meta.get("max_col", 0)

    # 行・列上限を適用
    effective_max_row = min(meta_max_row, max_rows) if meta_max_row > 0 else max_rows
    effective_max_col = min(meta_max_col, max_cols) if meta_max_col > 0 else max_cols

    # read_only モードでのセル読み込み
    # read_only モードでは iter_rows が軽量なジェネレータを返す
    for row_cells in ws.iter_rows(
        min_row=1,
        max_row=effective_max_row,
        min_col=1,
        max_col=effective_max_col,
    ):
        for opx_cell in row_cells:
            if opx_cell.value is None:
                # read_only モードではフォント情報が限定的なため、
                # 値がないセルはスキップ（結合セルは後でメタデータから補完）
                continue
            cell = _convert_cell_readonly(opx_cell)
            if cell is not None:
                cells[(cell.pos.row, cell.pos.col)] = cell

    # メタデータから結合セル情報を反映
    merged_cells = meta.get("merged_cells", [])
    merge_masters = meta.get("merge_masters", {})

    for (r, c), master_pos in merge_masters.items():
        if (r, c) in cells:
            cells[(r, c)].is_merged = True
            cells[(r, c)].merge_master = master_pos
        else:
            # 結合領域内の空セルも登録
            cells[(r, c)] = Cell(
                pos=CellPosition(row=r, col=c),
                is_merged=True,
                merge_master=master_pos,
            )

    return Sheet(
        name=ws.title,
        cells=cells,
        merged_cells=merged_cells,
        max_row=meta_max_row,
        max_col=meta_max_col,
        hidden=meta.get("sheet_state") == "hidden",
        hidden_rows=meta.get("hidden_rows", []),
        hidden_cols=meta.get("hidden_cols", []),
        row_groups=meta.get("row_groups", []),
        col_groups=meta.get("col_groups", []),
    )


def _convert_cell_readonly(opx_cell) -> Cell | None:
    """read_only モードのセルを内部 Cell に変換する。

    read_only モードではフォント・書式情報が限定的なため、
    値と位置情報のみ取得する。
    """
    row: int = opx_cell.row
    col: int = opx_cell.column

    value = opx_cell.value
    formula: str | None = None

    if isinstance(value, str) and value.startswith("="):
        formula = value
        value = None

    # read_only モードでも基本的な書式は取得可能
    fmt = CellFormat()

    # スタイル情報が取得可能な場合のみ設定
    try:
        if opx_cell.number_format and opx_cell.number_format != "General":
            fmt.number_format = opx_cell.number_format
    except (AttributeError, TypeError):
        pass

    return Cell(
        pos=CellPosition(row=row, col=col),
        value=value,
        fmt=fmt,
        formula=formula,
    )


def _parse_worksheet(
    ws: Worksheet,
    *,
    max_rows: int = DEFAULT_MAX_ROWS,
    max_cols: int = DEFAULT_MAX_COLS,
) -> Sheet:
    """Worksheet を Sheet モデルに変換する。"""
    cells: dict[tuple[int, int], Cell] = {}

    # 実際にデータが存在する範囲を特定（openpyxlのmax_columnは信頼できない場合がある）
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0

    # 行・列上限を適用
    effective_max_row = min(max_row, max_rows)

    # max_colが256(=IV列)の場合、実際の使用範囲を検出
    effective_max_col = min(max_col, max_cols)
    if max_col >= 256:
        effective_max_col = min(_detect_effective_max_col(ws, effective_max_row), max_cols)

    # セルの読み込み（値があるセルのみ保持）
    for row_cells in ws.iter_rows(min_row=1, max_row=effective_max_row, min_col=1, max_col=effective_max_col):
        for opx_cell in row_cells:
            if opx_cell.value is None and not (opx_cell.font and opx_cell.font.bold):
                continue
            cell = _convert_cell(opx_cell)
            if cell is not None:
                cells[(cell.pos.row, cell.pos.col)] = cell

    # 結合セル
    merged_cells: list[MergedRange] = []
    merge_masters: dict[tuple[int, int], CellPosition] = {}

    for merged_range in ws.merged_cells.ranges:
        mr = MergedRange(
            start_row=merged_range.min_row,
            start_col=merged_range.min_col,
            end_row=merged_range.max_row,
            end_col=merged_range.max_col,
        )
        merged_cells.append(mr)

        master_pos = CellPosition(row=mr.start_row, col=mr.start_col)
        for r in range(mr.start_row, mr.end_row + 1):
            for c in range(mr.start_col, mr.end_col + 1):
                merge_masters[(r, c)] = master_pos

    # 結合情報をセルに反映
    for (r, c), master_pos in merge_masters.items():
        if (r, c) in cells:
            cells[(r, c)].is_merged = True
            cells[(r, c)].merge_master = master_pos
        else:
            # 結合領域内の空セルも登録
            cells[(r, c)] = Cell(
                pos=CellPosition(row=r, col=c),
                is_merged=True,
                merge_master=master_pos,
            )

    # 非表示行
    hidden_rows: list[int] = []
    if ws.row_dimensions:
        for row_idx, rd in ws.row_dimensions.items():
            if rd.hidden:
                hidden_rows.append(int(row_idx))

    # 非表示列
    hidden_cols: list[int] = []
    if ws.column_dimensions:
        for col_key, cd in ws.column_dimensions.items():
            if cd.hidden:
                if isinstance(col_key, str):
                    col_idx = openpyxl.utils.column_index_from_string(col_key)
                else:
                    col_idx = int(col_key)
                hidden_cols.append(col_idx)

    # 行グループ / 列グループ
    row_groups = _extract_row_groups(ws)
    col_groups = _extract_col_groups(ws)

    return Sheet(
        name=ws.title,
        cells=cells,
        merged_cells=merged_cells,
        max_row=max_row,
        max_col=max_col,
        hidden=ws.sheet_state == "hidden",
        hidden_rows=sorted(hidden_rows),
        hidden_cols=sorted(hidden_cols),
        row_groups=row_groups,
        col_groups=col_groups,
    )


def _detect_effective_max_col(ws: Worksheet, max_row: int) -> int:
    """実際にデータが存在する最大列を検出する。

    openpyxl の max_column が 256 等の大きな値を返す場合に使用。
    先頭100行をサンプリングして実際の使用範囲を推定する。
    """
    sample_rows = min(max_row, 100)
    effective = 1
    for row_cells in ws.iter_rows(min_row=1, max_row=sample_rows):
        for cell in reversed(row_cells):
            if cell.value is not None:
                effective = max(effective, cell.column)
                break
    # 安全マージン（ヘッダーより下にデータが広がっている可能性）
    return min(effective + 5, ws.max_column or effective + 5)


def _convert_cell(opx_cell: OpenpyxlCell) -> Cell | None:
    """openpyxl セルを内部 Cell に変換する。"""
    row: int = opx_cell.row
    col: int = opx_cell.column

    value = opx_cell.value
    formula: str | None = None

    # data_only=True なので数式は取得できないが、
    # 型が str で '=' から始まる場合は数式の可能性
    if isinstance(value, str) and value.startswith("="):
        formula = value
        value = None

    # 書式情報
    font = opx_cell.font
    fill = opx_cell.fill

    font_color: str | None = None
    if font and font.color and font.color.rgb and font.color.rgb != "00000000":
        rgb = font.color.rgb
        if isinstance(rgb, str):
            font_color = rgb

    bg_color: str | None = None
    if fill and fill.fgColor and fill.fgColor.rgb and fill.fgColor.rgb != "00000000":
        rgb = fill.fgColor.rgb
        if isinstance(rgb, str):
            bg_color = rgb

    alignment_str: str | None = None
    if opx_cell.alignment:
        parts: list[str] = []
        if opx_cell.alignment.horizontal:
            parts.append(f"h:{opx_cell.alignment.horizontal}")
        if opx_cell.alignment.vertical:
            parts.append(f"v:{opx_cell.alignment.vertical}")
        if parts:
            alignment_str = ",".join(parts)

    border = opx_cell.border
    border_left = border.left.style if border else None
    border_right = border.right.style if border else None
    border_top = border.top.style if border else None
    border_bottom = border.bottom.style if border else None

    fmt = CellFormat(
        bold=bool(font and font.bold),
        italic=bool(font and font.italic),
        font_color=font_color,
        bg_color=bg_color,
        number_format=opx_cell.number_format if opx_cell.number_format != "General" else None,
        alignment=alignment_str,
        border_left=border_left,
        border_right=border_right,
        border_top=border_top,
        border_bottom=border_bottom,
    )

    return Cell(
        pos=CellPosition(row=row, col=col),
        value=value,
        fmt=fmt,
        formula=formula,
    )


def _extract_row_groups(ws: Worksheet) -> list[tuple[int, int]]:
    """行のグループ化情報を抽出する。"""
    groups: list[tuple[int, int]] = []
    if not ws.row_dimensions:
        return groups

    current_start: int | None = None
    for row_idx in sorted(ws.row_dimensions.keys()):
        rd = ws.row_dimensions[row_idx]
        if rd.outline_level and rd.outline_level > 0:
            if current_start is None:
                current_start = int(row_idx)
        else:
            if current_start is not None:
                groups.append((current_start, int(row_idx) - 1))
                current_start = None
    if current_start is not None:
        groups.append((current_start, max(int(k) for k in ws.row_dimensions.keys())))
    return groups


def _extract_col_groups(ws: Worksheet) -> list[tuple[int, int]]:
    """列のグループ化情報を抽出する。"""
    groups: list[tuple[int, int]] = []
    if not ws.column_dimensions:
        return groups

    col_indices: list[int] = []
    for col_key in ws.column_dimensions:
        if isinstance(col_key, str):
            col_indices.append(openpyxl.utils.column_index_from_string(col_key))
        else:
            col_indices.append(int(col_key))

    col_map: dict[int, Any] = {}
    for col_key in ws.column_dimensions:
        if isinstance(col_key, str):
            idx = openpyxl.utils.column_index_from_string(col_key)
        else:
            idx = int(col_key)
        col_map[idx] = ws.column_dimensions[col_key]

    current_start: int | None = None
    for col_idx in sorted(col_map.keys()):
        cd = col_map[col_idx]
        if cd.outline_level and cd.outline_level > 0:
            if current_start is None:
                current_start = col_idx
        else:
            if current_start is not None:
                groups.append((current_start, col_idx - 1))
                current_start = None
    if current_start is not None and col_map:
        groups.append((current_start, max(col_map.keys())))
    return groups


# ─── CSV / TSV ───


@dataclass
class _RawField:
    raw: list[str] = field(default_factory=list)
    value: list[str] = field(default_factory=list)
    was_quoted: bool = False
    malformed: bool = False
    has_unquoted_quote: bool = False

    @property
    def raw_text(self) -> str:
        return "".join(self.raw)

    @property
    def value_text(self) -> str:
        return "".join(self.value)


def _read_csv_bytes(file_name: str, file_bytes: IO[bytes], delimiter: str) -> WorkBook:
    """バイトストリームから CSV / TSV を読み込む。"""
    return _read_csv_data(file_name, file_bytes.read(), delimiter)


def _read_csv_data(file_name: str, raw: bytes, delimiter: str) -> WorkBook:
    """バイト列から CSV / TSV を読み込む。"""
    file_format = "csv" if delimiter == "," else "tsv"

    # エンコーディング判定
    has_bom = raw[:3] == b"\xef\xbb\xbf"
    try:
        encoding = "utf-8-sig" if has_bom else "utf-8"
        text = raw.decode(encoding)
    except UnicodeDecodeError:
        encoding = "cp932"
        text = raw.decode(encoding)

    sheet = _parse_csv_text(text, Path(file_name).stem, delimiter)
    return WorkBook(
        file_name=file_name,
        file_format=file_format,
        sheets=[sheet],
        file_size=len(raw),
        encoding=encoding,
    )


def _read_csv(path: Path, delimiter: str) -> WorkBook:
    """CSV / TSV を読み込む。"""
    file_format = "csv" if delimiter == "," else "tsv"

    # エンコーディングの自動検出 (UTF-8 → cp932 フォールバック)
    encoding = _detect_encoding(path)

    cells: dict[tuple[int, int], Cell] = {}
    max_row = 0
    max_col = 0

    with open(path, encoding=encoding, newline="") as f:
        reader = csv.reader(f, delimiter=delimiter)
        for row_idx, row_data in enumerate(reader, start=1):
            max_row = row_idx
            for col_idx, value in enumerate(row_data, start=1):
                if col_idx > max_col:
                    max_col = col_idx
                cell_value: Any = value if value != "" else None
                cells[(row_idx, col_idx)] = Cell(
                    pos=CellPosition(row=row_idx, col=col_idx),
                    value=cell_value,
                )

    with open(path, encoding=encoding, newline="") as f:
        diagnostics = _build_csv_diagnostics_from_lines(f, delimiter)

    sheet = Sheet(
        name=path.stem,
        cells=cells,
        max_row=max_row,
        max_col=max_col,
        csv_diagnostics=diagnostics,
    )
    return WorkBook(
        file_name=path.name,
        file_format=file_format,
        sheets=[sheet],
        file_size=path.stat().st_size,
        encoding=encoding,
    )


def _parse_csv_text(text: str, sheet_name: str, delimiter: str) -> Sheet:
    """CSV/TSVテキストからSheetを作る。値はcsv.reader、診断はraw scannerで取得する。"""
    cells: dict[tuple[int, int], Cell] = {}
    max_row = 0
    max_col = 0

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    for row_idx, row_data in enumerate(reader, start=1):
        max_row = row_idx
        for col_idx, value in enumerate(row_data, start=1):
            if col_idx > max_col:
                max_col = col_idx
            cell_value: Any = value if value != "" else None
            cells[(row_idx, col_idx)] = Cell(
                pos=CellPosition(row=row_idx, col=col_idx),
                value=cell_value,
            )

    diagnostics = _build_csv_diagnostics(text, delimiter)
    return Sheet(
        name=sheet_name,
        cells=cells,
        max_row=max_row,
        max_col=max_col,
        csv_diagnostics=diagnostics,
    )


def _build_csv_diagnostics(text: str, delimiter: str) -> CsvDiagnostics:
    return _build_csv_diagnostics_from_lines([text], delimiter)


def _build_csv_diagnostics_from_lines(lines, delimiter: str) -> CsvDiagnostics:
    """CSV raw diagnosticsを作る。

    全フィールドの詳細は保持せず、集計値と代表sampleだけを保存する。
    """
    diagnostics = CsvDiagnostics(
        delimiter=delimiter,
        max_issue_samples=MAX_CSV_ISSUE_SAMPLES,
        max_issues_before_abort=MAX_CSV_ISSUES_BEFORE_ABORT,
    )
    sample_counts: Counter[str] = Counter()
    field_count_histogram: Counter[int] = Counter()
    issue_total = 0

    logical_row = 1
    physical_line = 1
    record_start_line = 1
    record_raw: list[str] = []
    fields: list[_RawField] = []
    current = _RawField()
    at_field_start = True
    in_quotes = False
    after_quote = False
    abort_detailed_scan = False

    def add_issue(sample: CsvIssueSample) -> None:
        nonlocal issue_total, abort_detailed_scan
        diagnostics.issue_counts[sample.issue_type] = diagnostics.issue_counts.get(sample.issue_type, 0) + 1
        issue_total += 1
        if (
            len(diagnostics.issue_samples) < MAX_CSV_ISSUE_SAMPLES
            and sample_counts[sample.issue_type] < MAX_CSV_ISSUE_SAMPLES_PER_TYPE
        ):
            diagnostics.issue_samples.append(sample)
            sample_counts[sample.issue_type] += 1
        if issue_total > MAX_CSV_ISSUES_BEFORE_ABORT:
            diagnostics.truncated = True
            diagnostics.abort_reason = "too_many_issues"
            diagnostics.confidence = 0.85
            abort_detailed_scan = True

    def preview(value: str) -> str:
        return value[:CSV_RAW_PREVIEW_LIMIT]

    def cell_ref(row: int, col: int | None) -> str | None:
        if col is None:
            return None
        return str(CellPosition(row=row, col=col))

    def finish_field() -> None:
        nonlocal current, at_field_start, after_quote
        fields.append(current)
        current = _RawField()
        at_field_start = True
        after_quote = False

    def finish_record(parse_error: str | None = None) -> None:
        nonlocal fields, record_raw, logical_row, record_start_line
        if abort_detailed_scan:
            return

        finish_field()
        field_count = len(fields)
        field_count_histogram[field_count] += 1
        diagnostics.scanned_records = logical_row

        expected = diagnostics.expected_field_count
        if logical_row == 1 and field_count > 0:
            diagnostics.expected_field_count = field_count
            expected = field_count

        physical_end_line = physical_line
        if parse_error == "unclosed_quote" and record_raw and record_raw[-1] in ("\n", "\r"):
            physical_end_line = max(record_start_line, physical_line - 1)

        if parse_error == "unclosed_quote":
            add_issue(
                CsvIssueSample(
                    issue_type="unclosed_quote",
                    logical_row=logical_row,
                    physical_start_line=record_start_line,
                    physical_end_line=physical_end_line,
                    severity="error",
                    description="クォートが閉じられていません。",
                    raw_preview=preview("".join(record_raw)),
                    field_count=field_count,
                    expected_field_count=expected,
                )
            )
        elif physical_end_line > record_start_line:
            add_issue(
                CsvIssueSample(
                    issue_type="multiline_record",
                    logical_row=logical_row,
                    physical_start_line=record_start_line,
                    physical_end_line=physical_end_line,
                    severity="error",
                    description="論理行が複数の物理行にまたがっています。",
                    raw_preview=preview("".join(record_raw)),
                    field_count=field_count,
                    expected_field_count=expected,
                )
            )

        for col_idx, field_meta in enumerate(fields, start=1):
            value = field_meta.value_text
            if field_meta.has_unquoted_quote:
                add_issue(
                    CsvIssueSample(
                        issue_type="unquoted_quote",
                        logical_row=logical_row,
                        col=col_idx,
                        cell_range=cell_ref(logical_row, col_idx),
                        physical_start_line=record_start_line,
                        physical_end_line=physical_end_line,
                        severity="error",
                        description="ダブルクォーテーションを含むフィールドがクォートされていません。",
                        raw_preview=preview(field_meta.raw_text),
                        value_preview=preview(value),
                        was_quoted=field_meta.was_quoted,
                    )
                )

        if expected is not None and logical_row > 1 and field_count != expected:
            add_issue(
                CsvIssueSample(
                    issue_type="field_count_mismatch",
                    logical_row=logical_row,
                    physical_start_line=record_start_line,
                    physical_end_line=physical_end_line,
                    severity="warning",
                    description="基準列数と異なるため、未クォートの区切り文字で列数が崩れている可能性があります。",
                    raw_preview=preview("".join(record_raw)),
                    field_count=field_count,
                    expected_field_count=expected,
                )
            )

        logical_row += 1
        fields = []
        record_raw = []

    first_chunk = True
    for chunk in lines:
        if first_chunk:
            chunk = chunk.lstrip("\ufeff")
            first_chunk = False
        i = 0
        while i < len(chunk):
            ch = chunk[i]
            next_ch = chunk[i + 1] if i + 1 < len(chunk) else ""
            record_raw.append(ch)

            if abort_detailed_scan:
                if ch == "\n":
                    physical_line += 1
                elif ch == "\r":
                    physical_line += 1
                    if next_ch == "\n":
                        i += 1
                i += 1
                continue

            if in_quotes:
                current.raw.append(ch)
                if ch == '"':
                    if next_ch == '"':
                        current.raw.append(next_ch)
                        current.value.append('"')
                        record_raw.append(next_ch)
                        i += 1
                    else:
                        in_quotes = False
                        after_quote = True
                else:
                    current.value.append(ch)
                    if ch == "\n":
                        physical_line += 1
                    elif ch == "\r":
                        physical_line += 1
                        if next_ch == "\n":
                            current.raw.append(next_ch)
                            current.value.append(next_ch)
                            record_raw.append(next_ch)
                            i += 1
                i += 1
                continue

            if ch == delimiter:
                finish_field()
            elif ch in ("\n", "\r"):
                finish_record()
                physical_line += 1
                record_start_line = physical_line
                if ch == "\r" and next_ch == "\n":
                    i += 1
            elif ch == '"' and at_field_start:
                current.raw.append(ch)
                current.was_quoted = True
                in_quotes = True
                at_field_start = False
            elif ch == '"':
                current.raw.append(ch)
                current.value.append(ch)
                current.has_unquoted_quote = True
                at_field_start = False
            else:
                current.raw.append(ch)
                current.value.append(ch)
                if after_quote and not ch.isspace():
                    current.malformed = True
                    current.has_unquoted_quote = True
                at_field_start = False
            i += 1

    if in_quotes:
        finish_record(parse_error="unclosed_quote")
    elif current.raw or fields or record_raw:
        finish_record()

    diagnostics.total_records = logical_row - 1
    diagnostics.field_count_histogram = dict(field_count_histogram)
    if diagnostics.expected_field_count is None and field_count_histogram:
        diagnostics.expected_field_count = field_count_histogram.most_common(1)[0][0]
    return diagnostics


def _detect_encoding(path: Path) -> str:
    """UTF-8 で読めるか試し、ダメなら cp932 を返す。

    BOM 付き UTF-8 の場合は ``utf-8-sig`` を返す。
    ``utf-8-sig`` は BOM を自動的にスキップするため、
    csv.reader がフィールドのクォーティングを正しく解析できる。
    """
    try:
        raw_head = path.read_bytes()[:3]
        has_bom = raw_head[:3] == b"\xef\xbb\xbf"
        with open(path, encoding="utf-8") as f:
            f.read(4096)
        return "utf-8-sig" if has_bom else "utf-8"
    except UnicodeDecodeError:
        return "cp932"
